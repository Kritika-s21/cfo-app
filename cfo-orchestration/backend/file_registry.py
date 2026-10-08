"""Registry of data files that reached the EzCoworker workspace as input/<name>.

Files arrive from two places: the chat paperclip (source="chat") and the File Pickup Scheduler (source="scheduler").
Listing them lets the UI offer *every* uploaded file for any skill, whatever the file is called. Nothing here depends
on file naming conventions; column headers are stored only as a hint so users can tell files apart.
"""
import csv, datetime as _dt, io, json, os, re, threading, time
from pathlib import Path
from typing import Any, Dict, List, Optional

_lock = threading.Lock()


def sniff_columns(name: str, data: bytes, limit: int = 40) -> List[str]:
    """Best-effort header row of a csv/tsv/xlsx file. Never raises; returns [] when unknown."""
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    try:
        if ext in ("csv", "tsv", "txt"):
            lines = data[:65536].decode("utf-8-sig", errors="replace").splitlines()
            first = next((l for l in lines if l.strip()), "")
            if ext == "tsv":
                delim = "\t"
            else:
                try:
                    delim = csv.Sniffer().sniff(first, delimiters=",;\t|").delimiter
                except csv.Error:
                    delim = ","
            row = next(csv.reader([first], delimiter=delim), [])
            return [c.strip() for c in row if c.strip()][:limit]
        if ext in ("xlsx", "xlsm"):
            import openpyxl                       # optional dependency
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            for row in wb.worksheets[0].iter_rows(min_row=1, max_row=1, values_only=True):
                return [str(c).strip() for c in row if c not in (None, "")][:limit]
    except Exception:
        return []
    return []


_MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
_RE_MON_YEAR = re.compile(r"\b(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*[-\s,]+(\d{4})\b", re.I)
_RE_ISO = re.compile(r"\b(\d{4})[-/](\d{1,2})(?:[-/]\d{1,2})?\b")
_RE_DMY = re.compile(r"\b(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})\b")        # day-first (India); swapped when only month-first is valid
_DATE_HDR = re.compile(r"date|period|month|posting|invoice dt|\bdt\b|due", re.I)
_MAX_ROWS = 200_000


def _iter_rows(name: str, data: bytes):
    """Yield rows (lists of cell values) of the first sheet of a csv/tsv/xlsx file. Raises on unreadable data."""
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    if ext in ("csv", "tsv", "txt"):
        text = data.decode("utf-8-sig", errors="replace")
        first = next((l for l in text.splitlines() if l.strip()), "")
        if ext == "tsv":
            delim = "\t"
        else:
            try:
                delim = csv.Sniffer().sniff(first, delimiters=",;\t|").delimiter
            except csv.Error:
                delim = ","
        yield from csv.reader(io.StringIO(text), delimiter=delim)
    elif ext in ("xlsx", "xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for row in wb.worksheets[0].iter_rows(values_only=True):
            yield list(row)
    else:
        return


def _cell_months(cell) -> List[tuple]:
    """(year, month) pairs found in one cell."""
    out = []
    if isinstance(cell, (_dt.datetime, _dt.date)):
        out.append((cell.year, cell.month))
    elif isinstance(cell, str) and cell.strip():
        for m in _RE_MON_YEAR.finditer(cell):
            out.append((int(m.group(2)), _MON.index(m.group(1)[:3].title()) + 1))
        for m in _RE_ISO.finditer(cell):
            out.append((int(m.group(1)), int(m.group(2))))
        for m in _RE_DMY.finditer(cell):
            a, b, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
            out.append((y, b if b <= 12 else a))
    return [(y, mo) for y, mo in out if 2000 < y < 2100 and 1 <= mo <= 12]


def _num(v) -> float:
    try:
        return float(str(v).replace(",", "").replace("\u20b9", "").replace("$", "").strip())
    except ValueError:
        return 0.0


def sniff_date_range(name: str, data: bytes) -> Optional[Dict[str, Any]]:
    """Reporting period covered by a csv/tsv/xlsx file, in the same shape the browser used to compute for attached files:
    {year, startMonth, endMonth, months[], label, minYear, maxYear, totalCOGS, juneCOGS}. None when no dates are found.
    Columns whose header looks like a date/period are preferred, so invoice or PO numbers such as 2025-11 are not misread as months."""
    try:
        rows = _iter_rows(name, data)
        header = next(rows, None)
        if not header:
            return None
        hdr = [str(h or "").strip() for h in header]
        date_cols = {i for i, h in enumerate(hdr) if _DATE_HDR.search(h)}
        cogs_i = next((i for i, h in enumerate(hdr) if "cogs" in h.lower()), -1)
        mon_i = next((i for i, h in enumerate(hdr) if "month" in h.lower()), -1)
        found, total_cogs, june_cogs = set(), 0.0, 0.0
        for n, row in enumerate(rows):
            if n >= _MAX_ROWS:
                break
            for i, cell in enumerate(row):
                if date_cols and i not in date_cols:
                    continue
                found.update(_cell_months(cell))
            if cogs_i >= 0 and cogs_i < len(row):
                v = _num(row[cogs_i])
                if v > 0:
                    total_cogs += v
                    if mon_i >= 0 and mon_i < len(row) and _num(row[mon_i]) == 6:
                        june_cogs += v
        if not found:
            return None
        lo, hi = min(found), max(found)
        if lo[0] == hi[0]:
            label = f"{_MON[lo[1]-1]} {lo[0]}" if lo[1] == hi[1] else f"{_MON[lo[1]-1]}\u2013{_MON[hi[1]-1]} {lo[0]}"
        else:
            label = f"{_MON[lo[1]-1]} {lo[0]}\u2013{_MON[hi[1]-1]} {hi[0]}"
        return {"year": hi[0], "startMonth": lo[1], "endMonth": hi[1], "months": sorted({m for _, m in found}),
                "label": label, "minYear": lo[0], "maxYear": hi[0],
                "totalCOGS": total_cogs, "juneCOGS": june_cogs}
    except Exception:
        return None


class FileRegistry:
    def __init__(self, path: str = "files_registry.json"):
        self.path = Path(path)
        self.items: Dict[str, Dict[str, Any]] = {}
        if self.path.exists():
            try:
                self.items = json.loads(self.path.read_text(encoding="utf-8"))
            except (OSError, ValueError):
                self.items = {}

    def _save(self):
        tmp = self.path.with_suffix(".tmp")
        tmp.write_text(json.dumps(self.items, indent=1), encoding="utf-8")
        os.replace(tmp, self.path)

    def record(self, name: str, remote_path: str, size: int, source: str = "chat", source_ref: str = "",
               change: str = "", columns: Optional[List[str]] = None,
               date_range: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
        now = time.time()
        with _lock:
            old = self.items.get(name, {})
            item = {"name": name, "remote_path": remote_path, "size": size, "source": source or "chat",
                    "source_ref": source_ref, "change": change, "columns": columns or [], "date_range": date_range,
                    "first_seen": old.get("first_seen", now), "uploaded_at": now,
                    "versions": old.get("versions", 0) + 1}
            self.items[name] = item
            self._save()
            return item

    def list(self, source: str = "", q: str = "", limit: int = 200) -> List[Dict[str, Any]]:
        q = q.lower().strip()
        rows = [f for f in self.items.values()
                if (not source or f["source"] == source)
                and (not q or q in f["name"].lower() or any(q in c.lower() for c in f.get("columns", [])))]
        return sorted(rows, key=lambda f: f["uploaded_at"], reverse=True)[:limit]

    def remove(self, name: str) -> bool:
        """Forget the file in the picker. (The copy in the EzCoworker workspace is not touched.)"""
        with _lock:
            if name not in self.items:
                return False
            del self.items[name]
            self._save()
            return True
